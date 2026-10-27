# -*- coding: utf-8 -*-
"""
从一份真实排班模板生成「示例模板」：表头、列宽、字体、对齐全部保留，
只把数据行换成虚构样例。用于开源仓库——发布版不能带真实医院/科室/医生信息。

用法：
  python tools/make_sample_template.py <源模板.xlsx> <输出.xlsx>
"""
import sys

from openpyxl import load_workbook

SHEET = "号源排班"

# 导出表的 14 列（顺序与模板表头严格一致）
EXPECTED_HEADERS = [
    "号源日期", "班次", "班次开始时间", "班次结束时间", "机构名称", "坐诊科室名称",
    "业务科室名称", "坐诊医生编码", "号别名称", "诊室名称", "分诊台名称",
    "每个时间间隔号数", "时间间隔（分钟）", "现场号数",
]

# 虚构样例：日期用同一个月的连续几天，机构/科室/业务科室/医生/诊室/分诊台全部是占位名
SAMPLE_ROWS = [
    ["2027-03-01", "上午", "2027-03-01 08:00:00", "2027-03-01 12:00:00", "示例机构", "示例科室", "示例业务科室", "doctor001", "住院医师", "示例诊室1", "示例分诊台", "1", "8", "0"],
    ["2027-03-01", "下午", "2027-03-01 14:30:00", "2027-03-01 17:30:00", "示例机构", "示例科室", "示例业务科室", "doctor001", "住院医师", "示例诊室1", "示例分诊台", "1", "8", "0"],
    ["2027-03-01", "上午", "2027-03-01 08:00:00", "2027-03-01 12:00:00", "示例机构", "示例科室", "示例业务科室", "doctor002", "副主任医师", "示例诊室2", "示例分诊台", "1", "8", "0"],
    ["2027-03-01", "下午", "2027-03-01 14:30:00", "2027-03-01 17:30:00", "示例机构", "示例科室", "示例业务科室", "doctor002", "副主任医师", "示例诊室2", "示例分诊台", "1", "8", "0"],
    ["2027-03-02", "上午", "2027-03-02 08:00:00", "2027-03-02 12:00:00", "示例机构", "示例科室", "示例业务科室", "doctor003", "主治医师", "示例诊室3", "示例分诊台", "1", "8", "0"],
    ["2027-03-02", "下午", "2027-03-02 14:30:00", "2027-03-02 17:30:00", "示例机构", "示例科室", "示例业务科室", "doctor003", "主治医师", "示例诊室3", "示例分诊台", "1", "8", "0"],
]


def main(src, dst):
    wb = load_workbook(src)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb[wb.sheetnames[0]]

    # 校验源模板表头，避免源文件换了列顺序还按旧样例硬写
    headers = [ws.cell(row=1, column=c).value for c in range(1, len(EXPECTED_HEADERS) + 1)]
    actual = [str(h).strip() if h is not None else "" for h in headers]
    if actual != EXPECTED_HEADERS:
        raise SystemExit(f"源模板表头和预期不一致：\n  实际 {actual}\n  预期 {EXPECTED_HEADERS}")

    first_data_row = 2
    style_row = first_data_row
    last_row = ws.max_row
    max_col = ws.max_column

    # 清空原有数据（含样式来源行的值）
    for r in range(first_data_row, last_row + 1):
        for c in range(1, max_col + 1):
            ws.cell(row=r, column=c).value = None

    # 按样例回填，样式沿用原来第 2 行那一套
    for i, row in enumerate(SAMPLE_ROWS):
        r = first_data_row + i
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c)
            if i > 0:
                cell._style = ws.cell(row=style_row, column=c)._style
            cell.value = value
            # 整张模板的数据格都是文本格式，照抄（日期/数字都不会被 Excel 转掉）
            cell.number_format = "@"

    # 删掉多余的空白行，保持文件干净
    keep_last = first_data_row + len(SAMPLE_ROWS) - 1
    if last_row > keep_last:
        ws.delete_rows(keep_last + 1, last_row - keep_last)

    wb.save(dst)
    print(f"已写出示例模板：{dst}")
    print(f"  sheet={ws.title} 表头 {ws.max_column} 列，示例数据 {len(SAMPLE_ROWS)} 行")


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(2)
    main(sys.argv[1], sys.argv[2])
