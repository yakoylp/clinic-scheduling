# -*- coding: utf-8 -*-
"""
按「门诊排班模板.xlsx」的样式把排班数据写进 xlsx。

为什么不用 Node 的 xlsx 库直接写：SheetJS 重写工作表会丢掉模板的字体、居中、
列宽和文本格式，导出的表打开后观感会和模板差很远。openpyxl 改单元格时会保留
原有样式，所以这里用 Python 干这一步。

用法：python tools/xlsx_writer.py <job.json>
job.json:
{
  "template": "门诊排班模板.xlsx",   // 可为 null，表示新建工作簿
  "output":   "导出/号源排班-2026-10.xlsx",
  "sheet":    "号源排班",
  "headers":  [...10 列...],
  "rows":     [ [10 个字符串], ... ],
  "textColumns": [0,2,3,6]           // 需要按文本写入的列（避免被 Excel 当日期）
}
"""
import copy
import json
import re
import shutil
import sys

from openpyxl import Workbook, load_workbook

DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
DATETIME_RE = re.compile(r"^\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}$")
SKIP = {"!ref", "!cols", "!rows", "!merges", "!margins", "!freeze", "!autofilter"}


def is_datey(value):
    return bool(DATE_RE.match(value) or DATETIME_RE.match(value))


def copy_style(src, dst):
    """把源单元格的样式整套搬过去（字体/对齐/边框/填充/数字格式）。"""
    dst._style = copy.copy(src._style)
    dst.number_format = src.number_format
    dst.font = copy.copy(src.font)
    dst.alignment = copy.copy(src.alignment)
    dst.border = copy.copy(src.border)
    dst.fill = copy.copy(src.fill)
    dst.protection = copy.copy(src.protection)


def clear_merged(ws, first_data_row):
    """去掉数据区及其下方的合并单元格，模板里可能有跨行的合并。"""
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row >= first_data_row:
            ws.unmerge_cells(str(rng))


def find_first_data_row(ws):
    """第一行「A 列有值且不是表头」的行。模板里数据从第 2 行开始。"""
    for row in range(1, min(ws.max_row, 12) + 1):
        v = ws.cell(row=row, column=1).value
        if v is not None and not isinstance(v, str):
            return row
        if isinstance(v, str) and (DATE_RE.match(v.strip()) or DATETIME_RE.match(v.strip())):
            return row
    return 2


def main(job_path):
    with open(job_path, "r", encoding="utf-8") as fh:
        job = json.load(fh)

    template = job.get("template")
    output = job["output"]
    headers = job["headers"]
    rows = job["rows"]
    text_columns = set(job.get("textColumns") or [])

    if template:
        shutil.copyfile(template, output)
        wb = load_workbook(output)
    else:
        wb = Workbook()
        wb.remove(wb.active)

    sheet_name = job.get("sheet") or ("号源排班" if "号源排班" in wb.sheetnames else wb.sheetnames[0])
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.create_sheet(sheet_name)

    # 模板里的表头样式 / 数据行样式，取出来当基准。
    # 逐列取，因为模板表头各列样式并不完全一样（例如两段表头字体不同）。
    w = len(headers)
    header_styles = {}
    for c in range(w):
        cell = ws.cell(row=1, column=c + 1)
        if cell.value is not None or cell.has_style:
            header_styles[c] = copy.copy(cell._style)
    header_style_src = ws.cell(row=1, column=1)

    old_data_row = find_first_data_row(ws)
    data_style_src = ws.cell(row=old_data_row, column=1)
    fallback_styles = {}
    for c in range(w):
        cell = ws.cell(row=old_data_row, column=c + 1)
        if cell.value is not None or cell.has_style:
            fallback_styles[c] = cell

    clear_merged(ws, old_data_row)

    # 模板右侧（第 11 列起）常留有带样式的空格子，会让导出的表多出几列空白
    if ws.max_column > w:
        ws.delete_cols(w + 1, ws.max_column - w)

    max_needed_row = len(rows) + 1

    # 清掉模板里超出本次数据范围的数据行（含其样式残留）
    if ws.max_row > max_needed_row:
        ws.delete_rows(max_needed_row + 1, ws.max_row - max_needed_row)

    # 表头：逐列沿用模板里对应列的样式
    for c in range(w):
        cell = ws.cell(row=1, column=c + 1)
        cell.value = headers[c]
        if c in header_styles:
            cell._style = copy.copy(header_styles[c])
        elif header_style_src is not None:
            copy_style(header_style_src, cell)

    # 数据
    written = 0
    for i, data_row in enumerate(rows):
        r = i + 2
        for c in range(w):
            value = data_row[c] if c < len(data_row) else ""
            if value is None:
                value = ""
            value = str(value)
            if value.lower() == "none":
                value = ""
            cell = ws.cell(row=r, column=c + 1)
            cell.value = value if value != "" else None
            src = fallback_styles.get(c, data_style_src)
            if src is not None:
                copy_style(src, cell)
            if c in text_columns or is_datey(value):
                cell.number_format = "@"
        written += 1

    wb.save(output)
    ts = wb[sheet_name]
    print(
        json.dumps(
            {
                "ok": True,
                "output": output,
                "sheet": sheet_name,
                "rows": written,
                "columns": w,
                "maxRow": ts.max_row,
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    try:
        main(sys.argv[1])
    except Exception as exc:  # noqa: BLE001
        print(json.dumps({"ok": False, "error": f"{type(exc).__name__}: {exc}"}, ensure_ascii=False))
        sys.exit(1)
